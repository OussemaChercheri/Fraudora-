from datetime import date, datetime
from io import BytesIO
from typing import Optional
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.graphics.shapes import Drawing, Rect
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.lib.units import mm
from sqlalchemy import Date, and_, cast, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import joinedload

from app.models.duplicate_match import DuplicateMatch, DuplicateStatus, MatchType
from app.models.invoice import Invoice, InvoiceStatus
from app.models.invoice_data import InvoiceData
from app.models.user import User, UserRole
from app.services import analysis_service
from app.services.dashboard_service import get_dashboard_summary

PRIMARY = colors.HexColor("#185FA5")
SECONDARY = colors.HexColor("#1D9E75")
DANGER = colors.HexColor("#E24B4A")

FIELD_LABELS = {
    "invoice_number": "N° facture",
    "invoice_date": "Date facture",
    "supplier_name": "Fournisseur",
    "total_amount": "Montant TTC",
    "tax_amount": "TVA",
}


def _table_style(header_color=PRIMARY) -> TableStyle:
    return TableStyle(
        [
            ("BACKGROUND", (0, 0), (-1, 0), header_color),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
            (
                "ROWBACKGROUNDS",
                (0, 1),
                (-1, -1),
                [colors.white, colors.HexColor("#F5F7FA")],
            ),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]
    )


def _add_page_number(canvas, doc) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 9)
    canvas.setFillColor(colors.grey)
    canvas.drawCentredString(A4[0] / 2, 1 * cm, f"Page {doc.page}")
    canvas.restoreState()


def _missing_fields(data: InvoiceData) -> list[str]:
    missing: list[str] = []
    checks = [
        ("invoice_number", data.invoice_number, data.confidence_invoice_number),
        ("invoice_date", data.invoice_date, data.confidence_invoice_date),
        ("supplier_name", data.supplier_name, data.confidence_supplier_name),
        ("total_amount", data.total_amount, data.confidence_total_amount),
        ("tax_amount", data.tax_amount, data.confidence_tax_amount),
    ]
    for key, value, confidence in checks:
        if value is None or confidence < 0.5:
            missing.append(FIELD_LABELS[key])
    return missing


def _avg_confidence(data: InvoiceData) -> float:
    scores = [
        data.confidence_invoice_number,
        data.confidence_invoice_date,
        data.confidence_supplier_name,
        data.confidence_total_amount,
        data.confidence_tax_amount,
    ]
    return round(sum(scores) / len(scores), 4)


async def _get_review_required_rows(
    db: AsyncSession,
    user_id: Optional[UUID] = None,
) -> list[dict]:
    query = (
        select(Invoice, InvoiceData)
        .outerjoin(InvoiceData, InvoiceData.invoice_id == Invoice.id)
        .where(Invoice.status == InvoiceStatus.REVIEW_REQUIRED)
        .order_by(Invoice.created_at.desc())
    )
    if user_id is not None:
        query = query.where(Invoice.user_id == user_id)

    result = await db.execute(query)
    rows = []
    for invoice, data in result.all():
        missing = _missing_fields(data) if data else list(FIELD_LABELS.values())
        rows.append(
            {
                "filename": invoice.original_filename,
                "uploaded_at": invoice.created_at.strftime("%d/%m/%Y %H:%M"),
                "missing_fields": ", ".join(missing) if missing else "—",
                "action": "Vérification manuelle requise",
            }
        )
    return rows


async def _get_processed_invoice_rows(
    db: AsyncSession,
    user_id: Optional[UUID] = None,
) -> list[dict]:
    query = (
        select(Invoice, InvoiceData)
        .outerjoin(InvoiceData, InvoiceData.invoice_id == Invoice.id)
        .where(Invoice.status == InvoiceStatus.PROCESSED)
        .order_by(Invoice.created_at.desc())
    )
    if user_id is not None:
        query = query.where(Invoice.user_id == user_id)

    result = await db.execute(query)
    rows = []
    for invoice, data in result.all():
        rows.append(
            {
                "id": str(invoice.id),
                "filename": invoice.original_filename,
                "supplier": data.supplier_name if data else None,
                "invoice_number": data.invoice_number if data else None,
                "invoice_date": data.invoice_date.isoformat() if data and data.invoice_date else None,
                "total_amount": float(data.total_amount) if data and data.total_amount else None,
                "tax_amount": float(data.tax_amount) if data and data.tax_amount else None,
                "avg_confidence": _avg_confidence(data) if data else 0.0,
                "manually_corrected": data.is_manually_corrected if data else False,
            }
        )
    return rows


def _build_volume_chart(monthly: list[dict], width: float = 450, height: float = 200) -> Drawing:
    drawing = Drawing(width, height)
    if not monthly:
        return drawing

    max_val = max(max(item["count"], item["processed_count"]) for item in monthly) or 1
    bar_width = width / (len(monthly) * 3)
    chart_bottom = 30
    chart_height = height - 50

    for index, item in enumerate(monthly):
        x_base = 20 + index * (bar_width * 2.5)
        uploaded_height = (item["count"] / max_val) * chart_height
        processed_height = (item["processed_count"] / max_val) * chart_height

        drawing.add(
            Rect(
                x_base,
                chart_bottom,
                bar_width,
                uploaded_height,
                fillColor=PRIMARY,
                strokeColor=PRIMARY,
            )
        )
        drawing.add(
            Rect(
                x_base + bar_width + 4,
                chart_bottom,
                bar_width,
                processed_height,
                fillColor=SECONDARY,
                strokeColor=SECONDARY,
            )
        )

    return drawing


async def generate_pilot_report_pdf(
    db: AsyncSession,
    user_id: Optional[UUID] = None,
) -> bytes:
    summary = await analysis_service.get_processing_summary(db, user_id)
    ocr_stats = await analysis_service.get_ocr_quality_stats(db, user_id)
    suppliers = (await analysis_service.get_supplier_summary(db, user_id))[:20]
    monthly = await analysis_service.get_monthly_volume(db, months=6, user_id=user_id)
    review_rows = await _get_review_required_rows(db, user_id)

    buffer = BytesIO()
    report_date = datetime.now().strftime("%d/%m/%Y")
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CoverTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=24,
        textColor=PRIMARY,
        spaceAfter=12,
    )
    subtitle_style = ParagraphStyle(
        "CoverSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=14,
        textColor=colors.HexColor("#444444"),
        spaceAfter=8,
    )
    heading_style = ParagraphStyle(
        "SectionHeading",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=16,
        textColor=PRIMARY,
        spaceAfter=12,
    )

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=2 * cm,
        leftMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )
    story: list = []

    story.append(Spacer(1, 4 * cm))
    story.append(Paragraph("FraudGuard AI — Rapport d'Analyse", title_style))
    story.append(Paragraph("Résultats du traitement OCR", subtitle_style))
    story.append(Spacer(1, 1 * cm))
    story.append(Paragraph(f"Date : {report_date}", subtitle_style))
    story.append(Paragraph("Généré par : FraudGuard AI v1.0", subtitle_style))
    story.append(PageBreak())

    story.append(Paragraph("Résumé exécutif", heading_style))
    summary_table_data = [
        ["Indicateur", "Valeur"],
        ["Total factures", str(summary["total_invoices"])],
        ["Taux de traitement", f"{summary['processed_rate']}%"],
        ["Taux de révision requise", f"{summary['review_required_rate']}%"],
        ["Taux d'erreur", f"{summary['error_rate']}%"],
    ]
    summary_table = Table(summary_table_data, colWidths=[8 * cm, 8 * cm])
    summary_table.setStyle(_table_style())
    story.append(summary_table)
    story.append(Spacer(1, 0.5 * cm))

    confidence_table_data = [["Champ", "Confiance moyenne (%)"]]
    for field, value in ocr_stats["avg_confidence"].items():
        confidence_table_data.append(
            [FIELD_LABELS.get(field, field), f"{round(value * 100, 1)}"]
        )
    confidence_table = Table(confidence_table_data, colWidths=[8 * cm, 8 * cm])
    confidence_table.setStyle(_table_style(SECONDARY))
    story.append(confidence_table)
    story.append(PageBreak())

    story.append(Paragraph("Analyse par fournisseur", heading_style))
    supplier_table_data = [
        ["Fournisseur", "Nb factures", "Total montants", "Montant moyen"]
    ]
    for supplier in suppliers:
        supplier_table_data.append(
            [
                supplier["supplier_name"],
                str(supplier["invoice_count"]),
                f"{supplier['total_amount_sum']:.2f}",
                f"{supplier['avg_amount']:.2f}",
            ]
        )
    if len(supplier_table_data) == 1:
        supplier_table_data.append(["—", "0", "0.00", "0.00"])
    supplier_table = Table(
        supplier_table_data,
        colWidths=[6 * cm, 3 * cm, 4 * cm, 4 * cm],
    )
    supplier_table.setStyle(_table_style())
    story.append(supplier_table)
    story.append(PageBreak())

    story.append(Paragraph("Volume mensuel", heading_style))
    volume_table_data = [["Mois", "Uploaded", "Processed"]]
    for item in monthly:
        volume_table_data.append(
            [item["month"], str(item["count"]), str(item["processed_count"])]
        )
    volume_table = Table(volume_table_data, colWidths=[4 * cm, 4 * cm, 4 * cm])
    volume_table.setStyle(_table_style())
    story.append(volume_table)
    story.append(Spacer(1, 0.5 * cm))
    story.append(_build_volume_chart(monthly))
    story.append(PageBreak())

    story.append(Paragraph("Factures nécessitant révision", heading_style))
    review_table_data = [
        ["Fichier", "Date upload", "Champs manquants", "Action requise"]
    ]
    for row in review_rows:
        review_table_data.append(
            [
                row["filename"],
                row["uploaded_at"],
                row["missing_fields"],
                row["action"],
            ]
        )
    if len(review_table_data) == 1:
        review_table_data.append(["—", "—", "Aucune", "—"])
    review_table = Table(
        review_table_data,
        colWidths=[4.5 * cm, 3.5 * cm, 5 * cm, 4 * cm],
    )
    review_style = _table_style(DANGER)
    review_table.setStyle(review_style)
    story.append(review_table)

    doc.build(story, onFirstPage=_add_page_number, onLaterPages=_add_page_number)
    return buffer.getvalue()


def _auto_width(ws) -> None:
    for column_cells in ws.columns:
        max_length = 0
        column = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column].width = min(max_length + 2, 50)


async def generate_pilot_report_excel(
    db: AsyncSession,
    user_id: Optional[UUID] = None,
) -> bytes:
    summary = await analysis_service.get_processing_summary(db, user_id)
    ocr_stats = await analysis_service.get_ocr_quality_stats(db, user_id)
    suppliers = await analysis_service.get_supplier_summary(db, user_id)
    processed_rows = await _get_processed_invoice_rows(db, user_id)

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="185FA5")

    ws_summary = wb.active
    ws_summary.title = "Résumé"
    summary_rows = [
        ("Total factures", summary["total_invoices"]),
        ("Taux de traitement (%)", summary["processed_rate"]),
        ("Taux de révision requise (%)", summary["review_required_rate"]),
        ("Taux d'erreur (%)", summary["error_rate"]),
        ("Temps moyen traitement (ms)", summary["avg_processing_time_ms"]),
        ("Pages traitées", summary["total_pages_processed"]),
        ("Corrections manuelles", ocr_stats["manually_corrected_count"]),
        ("Taux corrections manuelles (%)", ocr_stats["manually_corrected_rate"]),
    ]
    ws_summary.append(["Indicateur", "Valeur"])
    for row in summary_rows:
        ws_summary.append(list(row))
    ws_summary.append([])
    ws_summary.append(["Confiance OCR moyenne par champ", ""])
    ws_summary.append(["Champ", "Confiance", "Taux extraction (%)"])
    for field, confidence in ocr_stats["avg_confidence"].items():
        ws_summary.append(
            [
                FIELD_LABELS.get(field, field),
                confidence,
                ocr_stats["field_extraction_rate"][field],
            ]
        )
    for cell in ws_summary[1]:
        cell.font = header_font
        cell.fill = header_fill

    ws_suppliers = wb.create_sheet("Fournisseurs")
    ws_suppliers.append(
        [
            "Fournisseur",
            "Nb factures",
            "Total montants",
            "Montant moyen",
            "Min",
            "Max",
            "Première facture",
            "Dernière facture",
        ]
    )
    for supplier in suppliers:
        ws_suppliers.append(
            [
                supplier["supplier_name"],
                supplier["invoice_count"],
                supplier["total_amount_sum"],
                supplier["avg_amount"],
                supplier["min_amount"],
                supplier["max_amount"],
                supplier["first_seen"].isoformat() if supplier["first_seen"] else "",
                supplier["last_seen"].isoformat() if supplier["last_seen"] else "",
            ]
        )
    for cell in ws_suppliers[1]:
        cell.font = header_font
        cell.fill = header_fill
    _auto_width(ws_suppliers)

    ws_invoices = wb.create_sheet("Factures")
    ws_invoices.append(
        [
            "ID",
            "Fichier",
            "Fournisseur",
            "N° facture",
            "Date",
            "Montant TTC",
            "TVA",
            "Confiance moy.",
            "Corrigé manuellement",
        ]
    )
    for row in processed_rows:
        ws_invoices.append(
            [
                row["id"],
                row["filename"],
                row["supplier"],
                row["invoice_number"],
                row["invoice_date"],
                row["total_amount"],
                row["tax_amount"],
                row["avg_confidence"],
                "Oui" if row["manually_corrected"] else "Non",
            ]
        )
    for cell in ws_invoices[1]:
        cell.font = header_font
        cell.fill = header_fill
    _auto_width(ws_invoices)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


async def generate_dashboard_report_pdf(
    db: AsyncSession,
    current_user: User,
    date_from: date,
    date_to: date,
) -> bytes:
    dash = await get_dashboard_summary(db, current_user, date_from=date_from, date_to=date_to)

    buffer = BytesIO()
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle("CoverTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=24, textColor=PRIMARY, spaceAfter=12)
    subtitle_style = ParagraphStyle("CoverSubtitle", parent=styles["Normal"], fontName="Helvetica", fontSize=14, textColor=colors.HexColor("#444444"), spaceAfter=8)
    heading_style = ParagraphStyle("SectionHeading", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=16, textColor=PRIMARY, spaceAfter=12)
    normal = ParagraphStyle("Normal", parent=styles["Normal"], fontName="Helvetica", fontSize=10)

    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=2*cm, leftMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm)
    story: list = []

    period_str = f"Du {date_from} au {date_to}"
    now_str = datetime.now().strftime("%d/%m/%Y %H:%M")

    story.append(Spacer(1, 4*cm))
    story.append(Paragraph("FraudGuard AI — Export Dashboard", title_style))
    story.append(Spacer(1, 1*cm))
    story.append(Paragraph(period_str, subtitle_style))
    story.append(Paragraph(f"Généré le {now_str}", subtitle_style))
    story.append(Spacer(1, 2*cm))
    status_summary = dash["factures_totales"]["by_status"]
    total_count = dash["factures_totales"]["total"]
    label_map = {"UPLOADED": "Importées", "PROCESSING": "En cours", "PROCESSED": "Traitées", "REVIEW_REQUIRED": "Révision", "ERROR": "Erreur"}
    status_lines = "<br/>".join(f"{label_map.get(k, k)}: {v}" for k, v in status_summary.items())
    story.append(Paragraph(f"<b>Total factures : {total_count}</b><br/>{status_lines}", normal))
    story.append(PageBreak())

    story.append(Paragraph("Factures totales", heading_style))
    status_data = [["Statut", "Nombre"]]
    status_data.extend([label_map.get(k, k), str(v)] for k, v in status_summary.items())
    status_data.append(["Total", str(total_count)])
    t = Table(status_data, colWidths=[8*cm, 4*cm])
    t.setStyle(_table_style())
    story.append(t)
    story.append(Spacer(1, 0.5*cm))

    story.append(Paragraph("Alertes en attente", heading_style))
    al = dash["alertes"]
    alert_data = [["Type", "Nombre"], ["Doublons", str(al["pending_duplicates"])], ["Anomalies", str(al["pending_anomalies"])], ["Total", str(al["total_pending"])]]
    t2 = Table(alert_data, colWidths=[8*cm, 4*cm])
    t2.setStyle(_table_style(SECONDARY))
    story.append(t2)
    story.append(Spacer(1, 0.5*cm))

    story.append(Paragraph("Factures suspectes", heading_style))
    suspect = dash["suspectes"]["count"]
    suspect_color = colors.HexColor("#FEF3C7") if suspect > 0 else colors.HexColor("#F5F7FA")
    suspect_text_color = colors.HexColor("#92400E") if suspect > 0 else colors.grey
    suspect_table = Table([["Factures suspectes", str(suspect)]], colWidths=[8*cm, 4*cm])
    suspect_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), suspect_color),
        ("TEXTCOLOR", (0, 0), (-1, -1), suspect_text_color),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 12),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#D97706") if suspect > 0 else colors.lightgrey),
    ]))
    story.append(suspect_table)
    story.append(PageBreak())

    story.append(Paragraph("Évolution", heading_style))
    evo = dash["evolution"]
    evo_data = [["Date/Semaine", "Importées", "Traitées"]]
    for i, lbl in enumerate(evo["labels"]):
        evo_data.append([lbl, str(evo["uploaded"][i]), str(evo["processed"][i])])
    t3 = Table(evo_data, colWidths=[5*cm, 4*cm, 4*cm])
    t3.setStyle(_table_style())
    story.append(t3)
    story.append(PageBreak())

    story.append(Paragraph("Top fournisseurs", heading_style))
    top = dash["top_fournisseurs"]["items"]
    top_data = [["Fournisseur", "Nb factures", "Montant total"]]
    for s in top:
        top_data.append([s["supplier_name"], str(s["invoice_count"]), f"{s['total_amount_sum']:.3f}"])
    if not top:
        top_data.append(["Aucun fournisseur", "0", "0"])
    t4 = Table(top_data, colWidths=[6*cm, 3*cm, 5*cm])
    t4.setStyle(_table_style())
    story.append(t4)
    story.append(Spacer(1, 0.5*cm))

    story.append(Paragraph("Scores de risque", heading_style))
    risk = dash["risques"]["items"]
    risk_data = [["Fournisseur", "Score"]]
    for s in risk:
        score = s["risk_score"]
        if score >= 60:
            row_color = colors.HexColor("#FEE2E2")
        elif score >= 30:
            row_color = colors.HexColor("#FFEDD5")
        else:
            row_color = colors.HexColor("#DCFCE7")
        risk_data.append([s["supplier_name"], str(score)])
    if not risk:
        risk_data.append(["Aucun fournisseur évalué", "—"])
    t5 = Table(risk_data, colWidths=[8*cm, 4*cm])
    t5.setStyle(_table_style())
    for i in range(1, len(risk_data)):
        score = risk[i-1]["risk_score"] if i-1 < len(risk) else 0
        if score >= 60:
            row_color = colors.HexColor("#FEE2E2")
        elif score >= 30:
            row_color = colors.HexColor("#FFEDD5")
        else:
            row_color = colors.HexColor("#DCFCE7")
        t5.setStyle(TableStyle([("BACKGROUND", (0, i), (-1, i), row_color)]))
    story.append(t5)

    doc.build(story, onFirstPage=_add_page_number, onLaterPages=_add_page_number)
    return buffer.getvalue()


async def generate_dashboard_report_excel(
    db: AsyncSession,
    current_user: User,
    date_from: date,
    date_to: date,
) -> bytes:
    dash = await get_dashboard_summary(db, current_user, date_from=date_from, date_to=date_to)

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="185FA5")

    ws_summary = wb.active
    ws_summary.title = "Résumé"
    ws_summary.append(["Indicateur", "Valeur"])
    label_map = {"UPLOADED": "Importées", "PROCESSING": "En cours", "PROCESSED": "Traitées", "REVIEW_REQUIRED": "Révision", "ERROR": "Erreur"}
    ws_summary.append(["Factures totales", dash["factures_totales"]["total"]])
    for k, v in dash["factures_totales"]["by_status"].items():
        ws_summary.append([f"  {label_map.get(k, k)}", v])
    ws_summary.append([])
    ws_summary.append(["Doublons en attente", dash["alertes"]["pending_duplicates"]])
    ws_summary.append(["Anomalies en attente", dash["alertes"]["pending_anomalies"]])
    ws_summary.append(["Total alertes", dash["alertes"]["total_pending"]])
    ws_summary.append([])
    ws_summary.append(["Factures suspectes", dash["suspectes"]["count"]])
    for cell in ws_summary[1]:
        cell.font = header_font
        cell.fill = header_fill
    _auto_width(ws_summary)

    ws_evo = wb.create_sheet("Évolution")
    ws_evo.append(["Date/Semaine", "Importées", "Traitées"])
    evo = dash["evolution"]
    for i, lbl in enumerate(evo["labels"]):
        ws_evo.append([lbl, evo["uploaded"][i], evo["processed"][i]])
    for cell in ws_evo[1]:
        cell.font = header_font
        cell.fill = header_fill
    _auto_width(ws_evo)

    ws_top = wb.create_sheet("Top fournisseurs")
    ws_top.append(["Fournisseur", "Nb factures", "Montant total"])
    for s in dash["top_fournisseurs"]["items"]:
        ws_top.append([s["supplier_name"], s["invoice_count"], s["total_amount_sum"]])
    for cell in ws_top[1]:
        cell.font = header_font
        cell.fill = header_fill
    _auto_width(ws_top)

    ws_risk = wb.create_sheet("Scores de risque")
    ws_risk.append(["Fournisseur", "Score de risque"])
    red_fill = PatternFill("solid", fgColor="FEE2E2")
    orange_fill = PatternFill("solid", fgColor="FFEDD5")
    green_fill = PatternFill("solid", fgColor="DCFCE7")
    for s in dash["risques"]["items"]:
        row_idx = ws_risk.max_row + 1
        ws_risk.append([s["supplier_name"], s["risk_score"]])
        score = s["risk_score"]
        if score >= 60:
            fill = red_fill
        elif score >= 30:
            fill = orange_fill
        else:
            fill = green_fill
        for col in range(1, 3):
            ws_risk.cell(row=row_idx, column=col).fill = fill
    for cell in ws_risk[1]:
        cell.font = header_font
        cell.fill = header_fill
    _auto_width(ws_risk)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


async def get_duplicates_report_data(
    db: AsyncSession,
    current_user: User,
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
) -> dict:
    if date_from is not None and isinstance(date_from, str):
        date_from = date.fromisoformat(date_from)
    if date_to is not None and isinstance(date_to, str):
        date_to = date.fromisoformat(date_to)

    conditions = [
        DuplicateMatch.status.in_([DuplicateStatus.PENDING, DuplicateStatus.CONFIRMED_DUPLICATE]),
    ]

    if current_user.role not in (UserRole.ADMIN, UserRole.FINANCE):
        owned = select(Invoice.id).where(Invoice.user_id == current_user.id)
        conditions.append(
            or_(
                DuplicateMatch.invoice_id.in_(owned),
                DuplicateMatch.matched_invoice_id.in_(owned),
            )
        )

    inv_filter = [Invoice.status != InvoiceStatus.ERROR]
    if date_from:
        inv_filter.append(cast(Invoice.created_at, Date) >= date_from)
    if date_to:
        inv_filter.append(cast(Invoice.created_at, Date) <= date_to)

    total_result = await db.execute(
        select(func.count()).select_from(Invoice).where(*inv_filter)
    )
    total_invoices = total_result.scalar() or 0

    query = select(DuplicateMatch).where(*conditions)
    if date_from or date_to:
        query = query.join(Invoice, DuplicateMatch.invoice_id == Invoice.id)
        if date_from:
            query = query.where(cast(Invoice.created_at, Date) >= date_from)
        if date_to:
            query = query.where(cast(Invoice.created_at, Date) <= date_to)

    query = query.order_by(DuplicateMatch.similarity_score.desc())
    result = await db.execute(query)
    matches = result.scalars().all()

    duplicate_count = len(matches)
    confirmed_count = sum(1 for m in matches if m.status == DuplicateStatus.CONFIRMED_DUPLICATE)
    pending_count = duplicate_count - confirmed_count
    duplicate_rate = round(duplicate_count / max(total_invoices, 1) * 100, 2)

    confirmed_ids = [m.matched_invoice_id for m in matches if m.status == DuplicateStatus.CONFIRMED_DUPLICATE]
    montant_a_risque_total = 0.0
    if confirmed_ids:
        amt_result = await db.execute(
            select(func.coalesce(func.sum(InvoiceData.total_amount), 0))
            .where(InvoiceData.invoice_id.in_(confirmed_ids))
        )
        montant_a_risque_total = float(amt_result.scalar() or 0)

    details = []
    for m in matches:
        inv_result = await db.execute(
            select(Invoice, InvoiceData)
            .outerjoin(InvoiceData, InvoiceData.invoice_id == Invoice.id)
            .where(Invoice.id == m.invoice_id)
        )
        inv_row = inv_result.first()
        inv, inv_data = inv_row if inv_row else (None, None)

        matched_result = await db.execute(
            select(Invoice, InvoiceData)
            .outerjoin(InvoiceData, InvoiceData.invoice_id == Invoice.id)
            .where(Invoice.id == m.matched_invoice_id)
        )
        mat_row = matched_result.first()
        mat, mat_data = mat_row if mat_row else (None, None)

        details.append({
            "invoice_filename": inv.original_filename if inv else "",
            "invoice_number": inv_data.invoice_number if inv_data else "",
            "supplier_name": inv_data.supplier_name if inv_data else "",
            "total_amount": float(inv_data.total_amount) if inv_data and inv_data.total_amount else 0.0,
            "invoice_date": inv_data.invoice_date.isoformat() if inv_data and inv_data.invoice_date else "",
            "matched_invoice_filename": mat.original_filename if mat else "",
            "matched_invoice_number": mat_data.invoice_number if mat_data else "",
            "match_type": m.match_type.value,
            "similarity_score": m.similarity_score,
            "status": m.status.value,
            "rejection_reason": m.rejection_reason or "",
        })

    return {
        "total_invoices_in_scope": total_invoices,
        "duplicate_count": duplicate_count,
        "confirmed_count": confirmed_count,
        "pending_count": pending_count,
        "duplicate_rate": duplicate_rate,
        "montant_a_risque_total": montant_a_risque_total,
        "details": details,
    }


async def generate_duplicates_report_pdf(
    db: AsyncSession,
    current_user: User,
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
) -> bytes:
    data = await get_duplicates_report_data(db, current_user, date_from, date_to)

    buffer = BytesIO()
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("CoverTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=24, textColor=PRIMARY, spaceAfter=12)
    subtitle_style = ParagraphStyle("CoverSubtitle", parent=styles["Normal"], fontName="Helvetica", fontSize=14, textColor=colors.HexColor("#444444"), spaceAfter=8)
    heading_style = ParagraphStyle("SectionHeading", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=16, textColor=PRIMARY, spaceAfter=12)
    normal = ParagraphStyle("Normal", parent=styles["Normal"], fontName="Helvetica", fontSize=10)

    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=2*cm, leftMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm)
    story: list = []

    period_str = f"Du {date_from}" if date_from else "Toutes les dates"
    if date_to:
        period_str += f" au {date_to}"
    now_str = datetime.now().strftime("%d/%m/%Y %H:%M")

    story.append(Spacer(1, 4*cm))
    story.append(Paragraph("FraudGuard AI — Rapport Doublons", title_style))
    story.append(Spacer(1, 1*cm))
    story.append(Paragraph(period_str, subtitle_style))
    story.append(Paragraph(f"Généré le {now_str}", subtitle_style))
    story.append(Spacer(1, 2*cm))

    summary_data = [
        ["Indicateur", "Valeur"],
        ["Total factures analysées", str(data["total_invoices_in_scope"])],
        ["Doublons détectés", str(data["duplicate_count"])],
        ["Confirmés", str(data["confirmed_count"])],
        ["En attente", str(data["pending_count"])],
        ["Taux de doublon", f"{data['duplicate_rate']}%"],
    ]
    t = Table(summary_data, colWidths=[8*cm, 6*cm])
    t.setStyle(_table_style())
    story.append(t)
    story.append(Spacer(1, 0.5*cm))

    risk_amount = data["montant_a_risque_total"]
    risk_box = Table(
        [["Montant total à risque (TND)", f"{risk_amount:,.3f}"]],
        colWidths=[8*cm, 6*cm],
    )
    if risk_amount > 0:
        risk_box.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FEE2E2")),
            ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#991B1B")),
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 12),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("BOX", (0, 0), (-1, -1), 2, DANGER),
        ]))
    else:
        risk_box.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#DCFCE7")),
            ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#166534")),
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 12),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("BOX", (0, 0), (-1, -1), 2, SECONDARY),
        ]))
    story.append(risk_box)
    story.append(PageBreak())

    story.append(Paragraph("Détail des doublons", heading_style))
    if not data["details"]:
        story.append(Paragraph("Aucun doublon détecté sur la période.", normal))
    else:
        detail_header = ["Facture", "N°", "Fournisseur", "Montant", "Type", "Simil.", "Statut"]
        col_widths = [3.5*cm, 2.5*cm, 3*cm, 2.5*cm, 2.5*cm, 1.5*cm, 2*cm]
        rows = [detail_header]
        for d in data["details"]:
            rows.append([
                d["invoice_filename"][:20],
                d["invoice_number"],
                d["supplier_name"][:15],
                f"{d['total_amount']:,.0f}",
                "Exact" if d["match_type"] == "EXACT_NUMBER" else "Flou",
                f"{d['similarity_score']:.0f}%",
                "Confirmé" if d["status"] == "CONFIRMED_DUPLICATE" else "En attente",
            ])

        per_page = 25
        for start in range(0, len(rows) - 1, per_page):
            if start > 0:
                story.append(PageBreak())
            end = min(start + per_page, len(rows))
            page_rows = [rows[0]] + rows[start + 1:end]
            t = Table(page_rows, colWidths=col_widths)
            base_style = _table_style()
            for i, d in enumerate(data["details"][start:end], 1):
                if d["status"] == "CONFIRMED_DUPLICATE":
                    base_style.add("BACKGROUND", (0, i), (-1, i), colors.HexColor("#FEE2E2"))
                elif d["status"] == "PENDING":
                    base_style.add("BACKGROUND", (0, i), (-1, i), colors.HexColor("#FEF3C7"))
            t.setStyle(base_style)
            story.append(t)

    doc.build(story, onFirstPage=_add_page_number, onLaterPages=_add_page_number)
    return buffer.getvalue()


async def generate_duplicates_report_excel(
    db: AsyncSession,
    current_user: User,
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
) -> bytes:
    data = await get_duplicates_report_data(db, current_user, date_from, date_to)

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="185FA5")

    ws = wb.active
    ws.title = "Résumé"
    ws.append(["Indicateur", "Valeur"])
    ws.append(["Total factures analysées", data["total_invoices_in_scope"]])
    ws.append(["Doublons détectés", data["duplicate_count"]])
    ws.append(["Confirmés", data["confirmed_count"]])
    ws.append(["En attente", data["pending_count"]])
    ws.append(["Taux de doublon (%)", data["duplicate_rate"]])
    ws.append(["Montant total à risque (TND)", data["montant_a_risque_total"]])
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
    _auto_width(ws)

    ws2 = wb.create_sheet("Détails")
    cols = ["Facture", "N° facture", "Fournisseur", "Montant", "Date", "Facture source", "N° source", "Type", "Similarité", "Statut", "Motif rejet"]
    ws2.append(cols)
    red_fill = PatternFill("solid", fgColor="FEE2E2")
    amber_fill = PatternFill("solid", fgColor="FEF3C7")

    for d in data["details"]:
        row_idx = ws2.max_row + 1
        ws2.append([
            d["invoice_filename"],
            d["invoice_number"],
            d["supplier_name"],
            d["total_amount"],
            d["invoice_date"],
            d["matched_invoice_filename"],
            d["matched_invoice_number"],
            "Exact" if d["match_type"] == "EXACT_NUMBER" else "Flou",
            d["similarity_score"],
            "Confirmé" if d["status"] == "CONFIRMED_DUPLICATE" else "En attente",
            d["rejection_reason"],
        ])
        if d["status"] == "CONFIRMED_DUPLICATE":
            for col in range(1, len(cols) + 1):
                ws2.cell(row=row_idx, column=col).fill = red_fill
        elif d["status"] == "PENDING":
            for col in range(1, len(cols) + 1):
                ws2.cell(row=row_idx, column=col).fill = amber_fill

    for cell in ws2[1]:
        cell.font = header_font
        cell.fill = header_fill
    _auto_width(ws2)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
